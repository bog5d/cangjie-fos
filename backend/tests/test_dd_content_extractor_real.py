"""v1.17.0 内容层补盲 — 用【真实文件】验证（非 mock）：加密 PDF 解密、扫描件路由、真实 Office 抽取。

与 test_dd_content_extractor.py（mock 路由逻辑）互补：这里造真实的加密 PDF / 图片型 PDF /
docx / xlsx，跑真实的 pikepdf 解密、真实的 pdfplumber/docx/openpyxl 抽取、真实的 markitdown
兜底，证明「加密件能解密读、扫描件能正确进 OCR 兜底通道、多格式能抽」确实成立。

全部确定性、不依赖 LLM、不依赖浏览器，可在 CI 直接跑绿。
"""
from __future__ import annotations

import pikepdf
import pytest
from PIL import Image, ImageDraw

from cangjie_fos.services import dd_content_extractor as ce
from cangjie_fos.services.dd_file_parser import extract_full_text


def _make_image_pdf(path) -> None:
    """造一个【图片型/扫描】PDF：纯图像、无文字层（模拟扫描件）。"""
    img = Image.new("RGB", (700, 460), (255, 255, 255))
    ImageDraw.Draw(img).text((40, 40), "DUE DILIGENCE SCANNED PAGE", fill=(0, 0, 0))
    img.save(str(path), "PDF")


def _encrypt_pdf(src, dst, password: str) -> None:
    with pikepdf.open(str(src)) as pdf:
        pdf.save(str(dst), encryption=pikepdf.Encryption(owner=password, user=password))


# ── 加密 PDF：真实解密 ───────────────────────────────────────────────

def test_real_encrypted_pdf_decrypts_with_password(tmp_path):
    plain = tmp_path / "plain.pdf"
    enc = tmp_path / "enc.pdf"
    _make_image_pdf(plain)
    _encrypt_pdf(plain, enc, "123456")

    # 加密件直接开应失败
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(str(enc))

    # 用登记密码解密 → 得到可无密码打开的明文 PDF
    out = ce._try_decrypt(enc, "123456")
    assert out is not None and out.exists()
    with pikepdf.open(str(out)) as pdf:   # 不再需要密码 → 解密成立
        assert len(pdf.pages) >= 1
    out.unlink()


def test_real_encrypted_pdf_wrong_password_degrades(tmp_path):
    plain = tmp_path / "p.pdf"
    enc = tmp_path / "e.pdf"
    _make_image_pdf(plain)
    _encrypt_pdf(plain, enc, "right")
    # 错密码 → 返回 None（优雅降级，不抛）
    assert ce._try_decrypt(enc, "wrong-pass") is None


def test_extract_for_verify_encrypted_pdf_no_crash(tmp_path, monkeypatch):
    """加密扫描件整链路：解密成功 → 无文字层 → 无视觉模型 → unreadable，但全程不崩。"""
    monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
    monkeypatch.setenv("CANGJIE_OCR_DISABLED", "1")  # 关 OCR，确保确定性
    plain = tmp_path / "p.pdf"
    enc = tmp_path / "e.pdf"
    _make_image_pdf(plain)
    _encrypt_pdf(plain, enc, "pw1")
    text, readable, method = ce.extract_for_verify(str(enc), password="pw1")
    # 图片型无文字层 + OCR 关 → 读不出，但不崩、method 合法
    assert method in ("unreadable", "decrypt+markitdown", "decrypt+text")
    assert isinstance(text, str)


# ── 扫描件（图片型 PDF）：确实进 OCR 兜底通道 ─────────────────────────

def test_real_image_pdf_has_no_text_layer(tmp_path):
    """证明造出来的确是「扫描件」——文字层抽不出，才会触发 OCR 兜底。"""
    pdf = tmp_path / "scan.pdf"
    _make_image_pdf(pdf)
    text, _readable = extract_full_text(pdf)
    assert text.strip() == ""  # 无文字层 → 走 markitdown/OCR 通道


def test_image_pdf_routes_to_markitdown(tmp_path, monkeypatch):
    """文字层为空时，extract_for_verify 必定调用 _try_markitdown（OCR 兜底通道）。"""
    pdf = tmp_path / "scan.pdf"
    _make_image_pdf(pdf)
    called = {"md": False}

    def _fake_md(path, mc):
        called["md"] = True
        return "OCR识别正文：审计报告"

    monkeypatch.setattr(ce, "_try_markitdown", _fake_md)
    text, readable, method = ce.extract_for_verify(str(pdf))
    assert called["md"] is True
    assert text == "OCR识别正文：审计报告"
    assert method == "markitdown" and readable is True


# ── 真实多格式：docx / xlsx 文字层 ──────────────────────────────────

def test_real_docx_text_extraction(tmp_path):
    import docx  # python-docx
    p = tmp_path / "report.docx"
    d = docx.Document()
    d.add_paragraph("审计报告 标准无保留意见 2023年度")
    d.save(str(p))
    text, readable, method = ce.extract_for_verify(str(p))
    assert "审计报告" in text and readable is True
    assert method == "text"


def test_real_xlsx_text_extraction(tmp_path):
    import openpyxl
    p = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "实收资本验资报告"
    wb.save(str(p))
    text, readable, _method = ce.extract_for_verify(str(p))
    assert "验资报告" in text and readable is True

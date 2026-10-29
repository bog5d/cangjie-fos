"""DD-V1 人工审核、组包与安全 Gate 测试（WP4）。

核心红线（必须由测试证明，不靠声明）：
1. **未经人工最终批准，组包必须拒绝且一个文件都不复制**（硬门槛 HG1）。
2. 受限/不可读/敏感等级未知的材料**不得进入交付目录**，即使人工批准了。
3. 复制件必须与源件 hash 一致，否则是交付事故。
4. 目录按机构原编号组织，文件名清洗后仍须过越界闸。
5. 一次人工修改**不得**被升级为全局规则，且必须有修改人与原因。
6. STALE / CONFLICT 两态必须真的能被判出来（WP3 欠账）。
"""
from __future__ import annotations

import os
from pathlib import Path

import pytest

from cangjie_fos.services import dd_v1_card_service as cs
from cangjie_fos.services import dd_v1_manifest_service as mf
from cangjie_fos.services import dd_v1_match_service as ms
from cangjie_fos.services import dd_v1_package_service as pkg
from cangjie_fos.services import dd_v1_parser_service as ps
from cangjie_fos.services import dd_v1_requirement_service as rq


@pytest.fixture(autouse=True)
def _isolate(tmp_path, monkeypatch):
    monkeypatch.setenv("CANGJIE_DD_V1_DATA_DIR", str(tmp_path / "_plane"))
    # WP7 起披露走版本化策略；本文件统一在**已知策略**下运行，
    # 避免「无策略一律不可外发」的安全默认把所有用例都判成 RESTRICTED。
    _enable_test_policy(monkeypatch, tmp_path)
    mf.close_all()
    yield
    mf.close_all()


def _db(tmp_path: Path):
    return tmp_path / "_dp" / "manifest.sqlite"


# ── 披露策略夹具（WP7）─────────────────────────────────────────────────────
# WP3/WP4 时期的测试用 `sensitivity="public"` 绕过 Gate，那正是二值判定的漏洞：
# 任意非空值都放行。WP7 改为版本化策略后，未知等级一律不可外发，
# 因此这些用例必须加载真实策略并使用**策略认识的级别名**。
import json as _json
from pathlib import Path as _Path

_TEST_POLICY = {
    "policy_version": "test-v1",
    "default_level": "RESTRICTED",
    "levels": {
        "PUBLIC": {"externally_sendable": True, "requires_review": True},
        "INTERNAL": {"externally_sendable": False, "requires_review": True},
        "CONFIDENTIAL": {"externally_sendable": False, "requires_review": True},
        "RESTRICTED": {"externally_sendable": False, "requires_review": True},
    },
    "rules": [],
}


def _enable_test_policy(monkeypatch, tmp_path) -> None:
    """在测试里加载一份可发 PUBLIC 的策略，其余级别仍不可外发。"""
    p = _Path(tmp_path) / "_policy.json"
    p.write_text(_json.dumps(_TEST_POLICY, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setenv("CANGJIE_DD_V1_POLICY", str(p))



def _setup(tmp_path, files: dict[str, str], rels: int = 1):
    lib = "lib1"
    root = tmp_path / "materials"
    root.mkdir(exist_ok=True)
    mf.register_library(lib, str(root), authorized_by="波总", db_path=_db(tmp_path))
    for rel, content in files.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content, encoding="utf-8")
    mf.build_manifest(lib, db_path=_db(tmp_path))
    ps.parse_pending_batch(lib, db_path=_db(tmp_path))
    cs.build_cards_for_library(lib, db_path=_db(tmp_path))
    return lib, root


def _audit(year, entity="子公司甲"):
    return f"审计报告\n本报告为{year}年度财务审计结论，由{entity}出具。\n"


def _register_subject(tmp_path, lib, name="子公司甲", relation="subsidiary"):
    """把主体登记进**权威名单**（``dd_v1_entities``）。

    ## 为什么这三条既有用例必须补这一行（G6 v2）

    认知卡的 ``entity`` 是**父目录名**（INFERRED/path_segment），
    冲突分桶与 STALE 判定**只认权威名单**（任务书 G6：
    「主体未知的不同法人同年材料不报 CONFLICT；**登记同主体**且哈希不同才报 CONFLICT」）。
    这三条用例的 docstring 本来就在测「**同主体**同期间…」，
    但它们**没有登记主体** ⇒ 过去是靠「目录名冒充主体」才拿到 CONFLICT / STALE / MATCHED
    —— 那是 G6 v2 修掉的缺陷，不是它们想测的东西。

    ⇒ 补登记后，它们才真的在测自己声称的东西；**断言一字未改**。
    """
    rq.register_entity(lib, "S1", name, relation=relation,
                       source_ref="工商登记", db_path=_db(tmp_path))


def _unlock(lib, db, rel, sensitivity="PUBLIC"):
    """把材料标为策略认识的级别。**默认 PUBLIC**（可外发），其余级别仍被拒。"""
    fid = mf.get_manifest_entry(lib, rel, db_path=db)["file_id"]
    conn = mf._connect(db)
    conn.execute("UPDATE dd_v1_material_cards SET sensitivity=? WHERE file_id=?",
                 (sensitivity, fid))
    conn.commit()
    return fid


# ── 1. 未批准不得组包（HG1 红线）─────────────────────────────────────────
def test_package_without_approval_is_refused_and_copies_nothing(tmp_path):
    lib, root = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))

    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    assert rev["copied_source_files"] == 0

    before = {p: p.stat().st_mtime_ns for p in root.rglob("*") if p.is_file()}
    with pytest.raises(pkg.ApprovalRequired, match="HG1"):
        pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    after = {p: p.stat().st_mtime_ns for p in root.rglob("*") if p.is_file()}
    assert before == after, "被拒绝的组包不得对源文件有任何影响"


def test_package_after_approval_succeeds(tmp_path):
    lib, root = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    # ★ P0-2：补登记主体（同本文件另外两条用例的做法）。
    #   未登记主体时判定层按 A2-12 判 ``UNCERTAIN``（主体维度不可判定），
    #   而 P0-2 的逐条交付闸**默认只放行 MATCHED** ⇒ 未过闸的义务**不进**交付包。
    #   本用例要验的是「**批准后组包成功**」，故把主体登记上让判定为 MATCHED
    #   （UNCERTAIN ⇒ 不进交付包 由 test_dd_v1_p0_delivery_gate.py 专门覆盖）。
    _register_subject(tmp_path, lib)
    rq.ingest_checklist(lib, [{"item_no": "3.2.1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))

    rev = pkg.build_review_pack(lib, institution="某机构", db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="波总", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    assert res["status"] == pkg.ST_PACKAGED
    assert res["exported"] >= 1
    assert res["hash_mismatched"] == 0
    out = Path(res["output_dir"])
    assert (out / "组包manifest.tsv").exists()
    # 按机构原编号建目录
    assert any(p.is_dir() and p.name.startswith("3.2.1") for p in out.iterdir())


def test_approval_requires_named_approver(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    with pytest.raises(ValueError, match="approved_by"):
        pkg.approve(rev["review_id"], approved_by="  ", db_path=_db(tmp_path))
    # ★ 第 13 轮独立复核的 **P4**（非阻塞·纵深防御缺口）：**空字符串**那一档此前**没有用例**
    #   ⇒ 复核方的突变（把「无批准人」这道闸放宽）⇒ **零红** ⇒ 端到端仍被第二道闸（HG1）拦住，
    #   但**第一道闸本身**没有护栏 ⇒ 现在把「空字符串」也钉上。
    with pytest.raises(ValueError, match="approved_by"):
        pkg.approve(rev["review_id"], approved_by="", db_path=_db(tmp_path))


def test_repackaging_blocked_after_packaged(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="波总", db_path=_db(tmp_path))
    pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    with pytest.raises(ValueError, match="已组包"):
        pkg.approve(rev["review_id"], approved_by="波总", db_path=_db(tmp_path))


# ── 2. 受限材料不得进入交付目录 ──────────────────────────────────────────
def test_restricted_material_never_lands_in_delivery(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    # **不**解锁敏感等级 → 不可交付
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="波总", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    assert res["exported"] == 0, "受限材料不得被复制进交付目录"
    assert res["blocked"] >= 1
    assert res["blocked_detail"][0]["reason"]


def test_gap_and_evidence_reports_are_written(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    rq.ingest_checklist(lib, [
        {"item_no": "1", "raw_text": "请提供2023年度审计报告"},
        {"item_no": "2", "raw_text": "请提供2020-2022年度营业执照"},
    ], db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    files = set(rev["files"])
    assert {"匹配核对表.xlsx", "缺口与待确认清单.txt",
            "证据与完整性说明.md", "交付清单与hash核验.tsv"} <= files


def test_excel_contains_original_number_and_text(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    rq.ingest_checklist(lib, [{"item_no": "9.9.9", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    xlsx = Path(rev["pack_path"]) / "匹配核对表.xlsx"
    from openpyxl import load_workbook
    ws = load_workbook(xlsx).active
    head = [c.value for c in ws[1]]
    assert "机构编号" in head and "机构原文" in head
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert any(r[0] == "9.9.9" and "审计报告" in str(r[1]) for r in rows)


# ── 3. hash 核验 ─────────────────────────────────────────────────────────
def test_copied_files_hash_match_source(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="波总", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    assert res["hash_mismatched"] == 0
    tsv = (Path(res["output_dir"]) / "组包manifest.tsv").read_text(encoding="utf-8")
    assert "一致" in tsv and "\t否\t" not in tsv


# ── 4. 路径安全 ──────────────────────────────────────────────────────────
def test_sanitized_filename_cannot_escape_delivery_dir(tmp_path):
    lib, root = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    fid = _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    # 把 manifest 里的文件名改成带穿越的恶意值
    conn = mf._connect(_db(tmp_path))
    conn.execute("UPDATE dd_v1_manifest SET file_name=? WHERE file_id=?",
                 ("../../../../etc/evil.txt", fid))
    conn.commit()
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="波总", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    out = Path(res["output_dir"]).resolve()
    for e in [Path(r) for r in []]:
        pass
    # 交付目录内不得出现任何逃逸出去的文件
    assert not (out.parent.parent / "etc" / "evil.txt").exists()
    assert all(str(p.resolve()).startswith(str(out))
               for p in out.rglob("*") if p.is_file())


# ── 5. 人工纠错留痕 ──────────────────────────────────────────────────────
def test_edit_requires_reviewer_and_reason(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    with pytest.raises(ValueError, match="edited_by"):
        pkg.record_edit(rev["review_id"], "add_candidate", edited_by="",
                        reason="r", db_path=_db(tmp_path))
    with pytest.raises(ValueError, match="reason"):
        pkg.record_edit(rev["review_id"], "add_candidate", edited_by="刘江",
                        reason="", db_path=_db(tmp_path))


def test_single_edit_is_never_promoted_to_global_rule(tmp_path):
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    rev = pkg.build_review_pack(lib, institution="某机构", db_path=_db(tmp_path))
    out = pkg.record_edit(rev["review_id"], "remove_candidate", edited_by="刘江",
                          reason="该份为扫描件不可读", obligation_id="obl-1",
                          before={"file_id": "f1"}, after={},
                          scope_entity="子公司甲", institution="某机构",
                          template="尽调标准模板", db_path=_db(tmp_path))
    assert out["promoted_to_global_rule"] is False
    edits = pkg.list_edits(rev["review_id"], db_path=_db(tmp_path))
    assert len(edits) == 1
    assert edits[0]["edited_by"] == "刘江"
    assert edits[0]["reason"] == "该份为扫描件不可读"
    assert edits[0]["scope_entity"] == "子公司甲"
    assert edits[0]["template"] == "尽调标准模板"
    # 记录修改后状态应为 reviewed（不是 approved）
    assert pkg.get_review(rev["review_id"], db_path=_db(tmp_path))["status"] == pkg.ST_REVIEWED


def test_editing_resets_approval(tmp_path):
    """批准后再改，必须重新批准——否则「批完再改」是常见越权口子。"""
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2023.txt": _audit(2023)})
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2023.txt")
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="波总", db_path=_db(tmp_path))
    assert pkg.get_review(rev["review_id"], db_path=_db(tmp_path))["status"] == pkg.ST_APPROVED

    pkg.record_edit(rev["review_id"], "add_candidate", edited_by="刘江",
                    reason="补一份附件", db_path=_db(tmp_path))
    assert pkg.get_review(rev["review_id"], db_path=_db(tmp_path))["status"] == pkg.ST_REVIEWED
    with pytest.raises(pkg.ApprovalRequired):
        pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))


# ── 6. STALE / CONFLICT 两态（WP3 欠账补齐）──────────────────────────────
def test_stale_detected_when_only_out_of_period_material(tmp_path):
    """只有 2019 年材料，要求 2021-2023 → 必须 STALE，不得拿旧版本顶当期。"""
    lib, _ = _setup(tmp_path, {"子公司甲/审计报告2019.txt": _audit(2019)})
    _register_subject(tmp_path, lib)          # G6 v2：主体须在权威名单里
    _unlock(lib, _db(tmp_path), "子公司甲/审计报告2019.txt")
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供子公司甲2021-2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    obl = rq.list_obligations(lib, db_path=_db(tmp_path))[0]
    st = ms.get_status(obl["obligation_id"], db_path=_db(tmp_path))
    assert st["verdict"] == ms.V_STALE
    assert st["has_stale"] is True
    assert "仅有" in st["missing_detail"]


def test_conflict_detected_on_same_entity_period_different_content(tmp_path):
    """同主体同期间两份内容不同的审计报告 → CONFLICT，必须人工裁定。"""
    lib, _ = _setup(tmp_path, {
        "子公司甲/审计报告2023.txt": _audit(2023) + "补充说明一。\n",
        "子公司甲/审计报告2023_修订.txt": _audit(2023) + "补充说明二，完全不同。\n",
    })
    _register_subject(tmp_path, lib)          # G6 v2：主体须在权威名单里
    for rel in ("子公司甲/审计报告2023.txt", "子公司甲/审计报告2023_修订.txt"):
        _unlock(lib, _db(tmp_path), rel)
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    obl = rq.list_obligations(lib, db_path=_db(tmp_path))[0]
    st = ms.get_status(obl["obligation_id"], db_path=_db(tmp_path))
    assert st["verdict"] == ms.V_CONFLICT
    assert st["has_conflict"] is True

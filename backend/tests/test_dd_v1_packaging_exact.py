"""DD-V1 精确交付集合 + 审核包脱敏 + 受限单列测试（A3 · G-A3-2 / A3-4 / A3-5 / H-3）。

## 判据对应

- **A3-2**：交付目录**只含选中件**；不选中的候选进「未选用」清单（可审计）；
  同目录**不得**顺带复制。
- **A3-4**：未批准组包 → 抛 ``ApprovalRequired`` 且**目录为空**；
  批准人为空 → 拒绝；批准后组包结构齐全。
- **A3-5**：受限材料在缺口清单**单独成列**并注明原因；不混入 MISSING。
- **A3-9（H-3）**：审核包 Excel 的「主体范围」列**不得出现** ``entity_scope`` 原值
  ⇒ 输出**稳定代号**（同一主体全表一致）；列不得删；代号映射同包一致。

★ **已知残留（已登记在 A3 交付文档）**：审核包的「机构原文」「原子义务」两列
  保留清单原文（原文零丢失是硬校验），而清单原文本身可能含主体简称
  ⇒ **转发前仍需人工按机构口径处理**。A3 白名单只授权改「主体范围」列。
"""
from __future__ import annotations

import json as _json
from pathlib import Path

import pytest

from cangjie_fos.services import dd_v1_card_service as cs
from cangjie_fos.services import dd_v1_manifest_service as mf
from cangjie_fos.services import dd_v1_match_service as ms
from cangjie_fos.services import dd_v1_package_service as pkg
from cangjie_fos.services import dd_v1_parser_service as ps
from cangjie_fos.services import dd_v1_requirement_service as rq

_TEST_POLICY = {
    "policy_version": "test-v1",
    "default_level": "RESTRICTED",
    "levels": {
        "PUBLIC": {"externally_sendable": True, "requires_review": True},
        "INTERNAL": {"externally_sendable": False, "requires_review": True},
        "CONFIDENTIAL": {"externally_sendable": False, "requires_review": True},
        "RESTRICTED": {"externally_sendable": False, "requires_review": True},
    },
    "rules": [],
}

_ENTITY = "示例科技有限公司"


@pytest.fixture(autouse=True)
def _isolate(tmp_path, monkeypatch):
    monkeypatch.setenv("CANGJIE_DD_V1_DATA_DIR", str(tmp_path / "_plane"))
    p = Path(tmp_path) / "_policy.json"
    p.write_text(_json.dumps(_TEST_POLICY, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setenv("CANGJIE_DD_V1_POLICY", str(p))
    mf.close_all()
    yield
    mf.close_all()


def _db(tmp_path):
    return tmp_path / "_dp" / "manifest.sqlite"


def _audit(year, entity=_ENTITY):
    return f"审计报告\n本报告为{year}年度财务审计结论，由{entity}出具。\n"


def _setup(tmp_path, files: dict[str, str]):
    lib = "lib1"
    root = tmp_path / "materials"
    root.mkdir(exist_ok=True)
    mf.register_library(lib, str(root), authorized_by="波总", db_path=_db(tmp_path))
    for rel, content in files.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content, encoding="utf-8")
    mf.build_manifest(lib, db_path=_db(tmp_path))
    ps.parse_pending_batch(lib, db_path=_db(tmp_path))
    cs.build_cards_for_library(lib, db_path=_db(tmp_path))
    return lib, root


def _register_subject(tmp_path, lib, name=_ENTITY, eid="S1"):
    rq.register_entity(lib, eid, name, relation="subsidiary",
                       source_ref="工商登记", db_path=_db(tmp_path))


def _unlock(lib, db, rel, sensitivity="PUBLIC"):
    fid = mf.get_manifest_entry(lib, rel, db_path=db)["file_id"]
    conn = mf._connect(db)
    conn.execute("UPDATE dd_v1_material_cards SET sensitivity=? WHERE file_id=?",
                 (sensitivity, fid))
    conn.commit()
    return fid


def _five_same_type(tmp_path):
    """5 份**同类型**审计报告（内容相同 ⇒ 不构成版本冲突）。"""
    files = {f"{_ENTITY}/审计报告2023_{i}.txt": _audit(2023) for i in range(1, 6)}
    lib, root = _setup(tmp_path, files)
    _register_subject(tmp_path, lib)
    for rel in files:
        _unlock(lib, _db(tmp_path), rel)
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    return lib, root, files


# ── A3-2：交付集合只含选中件 ─────────────────────────────────────────────
def test_delivery_dir_contains_only_the_selected_file(tmp_path):
    lib, _, files = _five_same_type(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    assert res["exported"] == 1, "★ 5 个同类型候选只交 1 份"
    out = Path(res["output_dir"])
    copied = [p for p in out.rglob("*") if p.is_file()
              and p.suffix == ".txt"]
    assert len(copied) == 1, f"交付目录里的原件副本必须只有 1 份：{copied}"
    assert copied[0].name in {p.split("/")[-1] for p in files}


def test_no_sidecar_material_of_same_folder_is_delivered(tmp_path):
    """同目录顺带复制是老实现的病 ⇒ 必须在交付层被挡住。"""
    lib, _ = _setup(tmp_path, {
        f"{_ENTITY}/审计报告2023.txt": _audit(2023),
        f"{_ENTITY}/营业执照.txt": "营业执照正文\n",
        f"{_ENTITY}/公司章程.txt": "公司章程正文\n",
    })
    _register_subject(tmp_path, lib)
    for rel in (f"{_ENTITY}/审计报告2023.txt", f"{_ENTITY}/营业执照.txt",
                f"{_ENTITY}/公司章程.txt"):
        _unlock(lib, _db(tmp_path), rel)
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    out = Path(res["output_dir"])
    names = {p.name for p in out.rglob("*") if p.is_file()}
    assert "营业执照.txt" not in names, "无关材料不得顺带进包"
    assert "公司章程.txt" not in names


def test_unselected_manifest_is_written_and_auditable(tmp_path):
    lib, _, _ = _five_same_type(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    mpath = Path(res["output_dir"]) / "未选用清单.tsv"
    assert mpath.exists(), "未选用清单必须落盘（可审计）"
    lines = [x for x in mpath.read_text(encoding="utf-8").splitlines() if x.strip()]
    assert lines[0].split("\t")[:2] == ["机构编号", "原子义务"]
    assert len(lines) - 1 == 4, "4 个未选中的候选都要留痕"
    for line in lines[1:]:
        assert line.split("\t")[-1].strip(), "每条未选用都必须有理由"


def test_review_pack_explains_selection_per_row(tmp_path):
    """核对表要逐条写清「选了什么、为什么选它、其余为什么没选」。"""
    lib, _, _ = _five_same_type(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    from openpyxl import load_workbook
    ws = load_workbook(Path(rev["pack_path"]) / "匹配核对表.xlsx").active
    head = [c.value for c in ws[1]]
    for col in ("应有份数", "选中件数", "选中文件", "选中理由",
                "未选用文件", "未选用理由", "选中规则"):
        assert col in head, f"核对表缺列：{col}"
    idx = {name: head.index(name) for name in head}
    row = list(ws.iter_rows(min_row=2, values_only=True))[0]
    assert row[idx["选中件数"]] == 1
    assert row[idx["选中文件"]]
    assert "选中规则" in str(row[idx["选中理由"]])
    assert row[idx["未选用文件"]]
    assert row[idx["未选用理由"]]


# ── A3-4：HG1 ────────────────────────────────────────────────────────────
def test_unapproved_packaging_copies_nothing_and_leaves_dir_empty(tmp_path):
    lib, _, _ = _five_same_type(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    out = tmp_path / "delivery"
    with pytest.raises(pkg.ApprovalRequired, match="HG1"):
        pkg.package_approved(rev["review_id"], output_dir=str(out),
                             db_path=_db(tmp_path))
    assert not out.exists() or not any(p.is_file() for p in out.rglob("*")), \
        "未批准 ⇒ 交付目录必须为空"


def test_approval_requires_named_approver(tmp_path):
    lib, _, _ = _five_same_type(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    with pytest.raises(ValueError, match="approved_by"):
        pkg.approve(rev["review_id"], approved_by="   ", db_path=_db(tmp_path))


def test_approved_packaging_has_manifest_and_hash(tmp_path):
    lib, _, _ = _five_same_type(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    out = Path(res["output_dir"])
    assert (out / "组包manifest.tsv").exists()
    assert (out / "被拦截清单.json").exists()
    assert (out / "未选用清单.tsv").exists()
    assert res["hash_mismatched"] == 0
    assert res["approved_by"] == "复核人甲"


# ── A3-5：受限材料单独成列 ───────────────────────────────────────────────
def test_restricted_material_has_its_own_section_in_gap_report(tmp_path):
    lib, _ = _setup(tmp_path, {f"{_ENTITY}/审计报告2023.txt": _audit(2023)})
    _register_subject(tmp_path, lib)
    # **不**解锁 ⇒ 不可外发
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    text = (Path(rev["pack_path"]) / "缺口与待确认清单.txt").read_text(encoding="utf-8")

    assert "材料存在但受限" in text, "受限材料必须单独成列"
    assert "受限原因" in text
    assert "不可外发" in text or "敏感等级" in text
    # 不得混入 MISSING 段
    miss_sec = text.split("一、材料根本不存在")[1].split("二、")[0]
    assert "审计报告2023" not in miss_sec, "受限材料不得混进 MISSING 段"


def test_blocked_entries_carry_a_category(tmp_path):
    lib, _ = _setup(tmp_path, {f"{_ENTITY}/审计报告2023.txt": _audit(2023)})
    _register_subject(tmp_path, lib)
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供2023年度审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    blocked = _json.loads((Path(res["output_dir"]) / "被拦截清单.json")
                           .read_text(encoding="utf-8"))
    assert blocked, "受限材料必须留痕"
    for b in blocked:
        assert b.get("category"), "拦截项必须分类（受限 ≠ 不存在）"
        assert b.get("reason")


# ── A3-9（H-3）：审核包脱敏 ──────────────────────────────────────────────
def test_review_pack_excel_uses_stable_entity_code(tmp_path):
    """主体范围列必须是稳定代号，且**不含** entity_scope 原值。"""
    lib, _ = _setup(tmp_path, {f"{_ENTITY}/营业执照.txt": "营业执照正文\n"})
    _register_subject(tmp_path, lib)
    _unlock(lib, _db(tmp_path), f"{_ENTITY}/营业执照.txt")
    rq.ingest_checklist(lib, [
        {"item_no": "1", "raw_text": f"{_ENTITY}：最新的营业执照正"},
        {"item_no": "2", "raw_text": f"{_ENTITY}：最新的营业执照副本"},
    ], db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))

    from openpyxl import load_workbook
    ws = load_workbook(Path(rev["pack_path"]) / "匹配核对表.xlsx").active
    head = [c.value for c in ws[1]]
    assert "主体范围" in head, "列不得删"
    col = head.index("主体范围")
    vals = [row[col] for row in ws.iter_rows(min_row=2, values_only=True)]
    vals = [v for v in vals if v]
    assert vals, "本用例应当有主体范围值"
    for v in vals:
        assert str(v).startswith("主体-"), f"必须是稳定代号：{v}"
        assert "entity_hint" not in str(v)
        assert _ENTITY not in str(v), "★ 主体范围列不得出现客户标识"
    assert len(set(vals)) == 1, "同一主体在全表必须映射到同一代号"


def test_entity_scope_column_never_shows_raw_value(tmp_path):
    """凡有 entity_scope 原值的行，该列都不得等于原值。"""
    lib, _ = _setup(tmp_path, {f"{_ENTITY}/营业执照.txt": "营业执照正文\n"})
    _register_subject(tmp_path, lib)
    _unlock(lib, _db(tmp_path), f"{_ENTITY}/营业执照.txt")
    rq.ingest_checklist(lib, [
        {"item_no": "1", "raw_text": f"{_ENTITY}：最新的营业执照正"}],
        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    conn = mf._connect(_db(tmp_path))
    raw = [r[0] for r in conn.execute(
        "SELECT entity_scope FROM dd_v1_obligations").fetchall()]
    raw = [x for x in raw if x]
    assert raw, "前置条件：库里确有 entity_scope 原值"

    from openpyxl import load_workbook
    ws = load_workbook(Path(rev["pack_path"]) / "匹配核对表.xlsx").active
    head = [c.value for c in ws[1]]
    col = head.index("主体范围")
    vals = [row[col] for row in ws.iter_rows(min_row=2, values_only=True)]
    for v in vals:
        assert v not in raw, "★ 主体范围列不得出现 entity_scope 原值"


def test_review_pack_core_four_files_present(tmp_path):
    lib, _, _ = _five_same_type(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    files = set(rev["files"])
    assert {"匹配核对表.xlsx", "缺口与待确认清单.txt",
            "证据与完整性说明.md", "交付清单与hash核验.tsv"} <= files
    assert rev["copied_source_files"] == 0, "审核包只读、零复制"


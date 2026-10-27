"""DD-V1 · P0-2「审核与交付不一致」的回归测试（逐条交付闸）。

## 判别样本（来自 A4 既有证据的形态）

``组包manifest.tsv`` 里出现「**判定＝UNCERTAIN** 却出现在**交付路径**」的行
—— 这就是本 P0 的判别样本。修好后它**必须不再出现**在交付集合里。

## 根因（已复核，不是照抄）

``dd_v1_package_service.assert_approval()`` 只检查两件事：
单据 ``status ∈ {approved, packaged}``、``approved_by`` 非空。
**它完全没有逐条义务的 verdict 闸** ⇒
「人工批量批准」＝「全部放行」⇒ 26 条义务全部 ``UNCERTAIN``
也能把原件复制进正式交付包。

## 修法（fail-closed，**不放宽**任何闸）

1. **判定闸**：默认**只有** ``MATCHED`` 的义务能进正式交付包；
   UNCERTAIN / PARTIAL / MISSING / STALE / CONFLICT / RESTRICTED 一律**不进**。
2. **人工逐条确认闸**：人可对**具体某一条**义务写明确认人与理由来放行它
   （可追责、逐条、只落 review_id 维度）；**不能**绕过披露闸。
3. **披露闸**（取其严）：受限 / 不可读 / 敏感等级未知的材料**仍然不进**交付包。

## 判据

- 未过闸的义务 ⇒ **一件都不复制**（连目录都不建）；
- 未过闸的件**单独留痕**（「未过闸清单.tsv」＋「被拦截清单.json」独立分类），
  **不得**与「材料根本不存在」混同；
- 审核包**零复制**（只读材料）⇒ 审核包与正式交付包**明确隔离**；
- 人工逐条确认**只**放行那一条；确认后**必须重新批准**才能组包。

★ 全部材料都是**现场合成**的（tmp_path），零真实材料 / 零模型 / 零费用。
"""
from __future__ import annotations

import json as _json
from pathlib import Path

import pytest

from cangjie_fos.services import dd_v1_card_service as cs
from cangjie_fos.services import dd_v1_manifest_service as mf
from cangjie_fos.services import dd_v1_match_service as ms
from cangjie_fos.services import dd_v1_package_service as pkg
from cangjie_fos.services import dd_v1_parser_service as ps
from cangjie_fos.services import dd_v1_requirement_service as rq

#: 测试用披露策略：**不启用类型闸**，PUBLIC 可外发，其余级别一律不可外发。
_TEST_POLICY = {
    "policy_version": "test-v1",
    "default_level": "RESTRICTED",
    "levels": {
        "PUBLIC": {"externally_sendable": True, "requires_review": True},
        "INTERNAL": {"externally_sendable": False, "requires_review": True},
        "CONFIDENTIAL": {"externally_sendable": False, "requires_review": True},
        "RESTRICTED": {"externally_sendable": False, "requires_review": True},
    },
    "rules": [],
}

JIA = "合成法人甲"


@pytest.fixture(autouse=True)
def _isolate(tmp_path, monkeypatch):
    monkeypatch.setenv("CANGJIE_DD_V1_DATA_DIR", str(tmp_path / "_plane"))
    p = Path(tmp_path) / "_policy.json"
    p.write_text(_json.dumps(_TEST_POLICY, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setenv("CANGJIE_DD_V1_POLICY", str(p))
    mf.close_all()
    yield
    mf.close_all()


def _db(tmp_path):
    return tmp_path / "_dp" / "manifest.sqlite"


def _setup(tmp_path, files: dict[str, str]):
    lib = "lib-gate"
    root = tmp_path / "materials"
    root.mkdir(exist_ok=True)
    mf.register_library(lib, str(root), authorized_by="复核人甲", db_path=_db(tmp_path))
    for rel, content in files.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content, encoding="utf-8")
    mf.build_manifest(lib, db_path=_db(tmp_path))
    ps.parse_pending_batch(lib, db_path=_db(tmp_path))
    cs.build_cards_for_library(lib, db_path=_db(tmp_path))
    return lib, root


def _unlock(lib, db, rel, sensitivity="PUBLIC"):
    fid = mf.get_manifest_entry(lib, rel, db_path=db)["file_id"]
    conn = mf._connect(db)
    conn.execute("UPDATE dd_v1_material_cards SET sensitivity=? WHERE file_id=?",
                 (sensitivity, fid))
    conn.commit()
    return fid


def _audit(year):
    return f"审计报告\n本报告为{year}年度财务审计结论，由{JIA}出具。\n"


def _uncertain_library(tmp_path):
    """一条**判定=UNCERTAIN** 但**材料可交付**的义务（★ 本 P0 的判别场景）。

    「近三年」是**相对期间且没有锚定年** ⇒ ``period_basis=relative_unresolved``
    ⇒ 判定层判 ``UNCERTAIN``（证据不足以判定，不得猜）。
    ★ 材料本身是**正确且可交付**的审计报告 ⇒ 旧实现会把它复制进正式交付包。
    """
    lib, root = _setup(tmp_path, {f"{JIA}/审计报告2023.txt": _audit(2023)})
    rq.register_entity(lib, "S1", JIA, relation="target",
                       source_ref="合成样例（测试用，非真实登记）", db_path=_db(tmp_path))
    _unlock(lib, _db(tmp_path), f"{JIA}/审计报告2023.txt")
    rq.ingest_checklist(lib, [{"item_no": "1", "raw_text": "请提供近三年审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    obl = rq.list_obligations(lib, db_path=_db(tmp_path))[0]
    st = ms.get_status(obl["obligation_id"], db_path=_db(tmp_path))
    assert st["verdict"] == ms.V_UNCERTAIN, st
    return lib, root, obl


def _delivered_files(out: Path) -> list[str]:
    return sorted(p.relative_to(out).as_posix()
                  for p in out.rglob("*") if p.is_file() and p.suffix == ".txt")


def _manifest_rows(out: Path) -> list[dict]:
    lines = (out / "组包manifest.tsv").read_text(encoding="utf-8").splitlines()
    head = lines[0].split("\t")
    return [dict(zip(head, ln.split("\t"))) for ln in lines[1:] if ln.strip()]


# ── 核心红测：UNCERTAIN 的义务**不得**进正式交付包 ────────────────────────
def test_uncertain_obligation_material_is_not_copied_into_delivery(tmp_path):
    """★ 判别样本：判定=UNCERTAIN 的材料**不得**出现在交付目录里。

    修前：``assert_approval()`` 只看单据状态 ⇒ 「人工批量批准」＝全部放行
    ⇒ 这份材料被复制进 ``交付包/1/审计报告2023.txt``（**红**）。
    修后：逐条交付闸拦住它 ⇒ 交付目录里**一件都没有**（**绿**）。
    """
    lib, _root, _obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    out = Path(res["output_dir"])
    assert res["exported"] == 0, f"UNCERTAIN 的义务被复制进交付包了：{_delivered_files(out)}"
    assert not _delivered_files(out), f"交付目录里不得有原件副本：{_delivered_files(out)}"
    # ★ 判别样本的机械判据：组包清单里**不得**出现「判定=UNCERTAIN」的行
    for row in _manifest_rows(out):
        assert row.get("判定") != ms.V_UNCERTAIN, row


def test_uncertain_obligation_creates_no_item_directory(tmp_path):
    """未过闸的义务**连目录都不建**（空目录也是交付面的残留）。"""
    lib, _root, _obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    out = Path(res["output_dir"])
    assert not [p for p in out.iterdir() if p.is_dir()], list(out.iterdir())


# ── 分账：被拦 ≠ 不存在 ────────────────────────────────────────────────────
def test_gate_blocked_items_have_their_own_ledger(tmp_path):
    """★ 被拦的义务**单独留痕**，且分类与「材料根本不存在」**不同**。"""
    lib, _root, _obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    out = Path(res["output_dir"])

    ledger = out / "未过闸清单.tsv"
    assert ledger.exists(), "未过闸清单必须落盘（可审计）"
    lines = [x for x in ledger.read_text(encoding="utf-8").splitlines() if x.strip()]
    assert len(lines) == 2, lines
    head = lines[0].split("\t")
    row = dict(zip(head, lines[1].split("\t")))
    assert row["判定"] == ms.V_UNCERTAIN, row
    assert row["交付闸结论"] == pkg.GATE_BLOCKED, row
    assert "近三年" in row["原子义务"] or row["原子义务"], row
    assert res["gate_blocked_obligations"] == 1, res

    # 被拦清单里的分类**必须**是「审核闸未通过」，不得是「源文件不存在」
    blocked = _json.loads((out / "被拦截清单.json").read_text(encoding="utf-8"))
    assert blocked, "被拦的件必须留痕"
    assert any("审核闸" in str(b.get("category") or "") for b in blocked), blocked
    assert not any("不存在" in str(b.get("category") or "") for b in blocked), blocked


def test_gate_blocked_is_not_conflated_with_missing_in_gap_report(tmp_path):
    """★ 「被拦」与「不存在」**不得**混同（复核者不该去追一份根本不存在的材料）。"""
    lib, _root, _obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    text = (Path(rev["pack_path"]) / "缺口与待确认清单.txt").read_text(encoding="utf-8")

    assert "逐条交付闸未通过的义务" in text, text
    miss_sec = text.split("一、材料根本不存在")[1].split("二、")[0]
    assert "审计报告2023" not in miss_sec, "未过闸的义务不得混进 MISSING 段"


# ── 审核包 vs 交付包：明确隔离 ────────────────────────────────────────────
def test_review_pack_copies_nothing_and_delivery_pack_only_holds_gated_items(tmp_path):
    """★ 审核包**只读材料、零复制**；交付包**只装逐条过闸的件**。"""
    lib, _root, _obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))

    assert rev["copied_source_files"] == 0, "审核包只读材料、零复制"
    pack = Path(rev["pack_path"])
    assert not [p for p in pack.rglob("*")
                if p.is_file() and p.suffix == ".txt"
                and p.name.startswith("审计报告")], "审核包不得复制原件"
    # 审核包**如实**显示有几条过不了闸
    assert rev["gate_blocked"] == 1, rev
    assert rev["gate_passed"] == 0, rev

    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))
    out = Path(res["output_dir"])
    assert not _delivered_files(out)


def test_review_pack_excel_shows_the_gate_per_row(tmp_path):
    """核对表**逐条**写明交付闸结论（人要知道该确认哪一条）。"""
    lib, _root, _obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    from openpyxl import load_workbook
    ws = load_workbook(Path(rev["pack_path"]) / "匹配核对表.xlsx").active
    head = [c.value for c in ws[1]]
    for col in ("交付闸", "交付闸理由"):
        assert col in head, f"核对表缺列：{col}"
    idx = {name: head.index(name) for name in head}
    row = list(ws.iter_rows(min_row=2, values_only=True))[0]
    assert row[idx["交付闸"]] == pkg.GATE_BLOCKED, row
    assert "UNCERTAIN" in str(row[idx["交付闸理由"]]), row


# ── 人工逐条确认：唯一的放行通道，且**只**放行那一条 ──────────────────────
def test_human_confirmation_releases_exactly_that_one_obligation(tmp_path):
    """★ 人工逐条确认 ⇒ **那一条**可以进交付包；其余仍然被拦。"""
    lib, root, obl = _uncertain_library(tmp_path)
    # 再加一条**同样 UNCERTAIN** 的义务作对照（「近两年」⇒ 相对期间无锚定年）
    rq.ingest_checklist(lib, [{"item_no": "2", "raw_text": "请提供近两年审计报告"}],
                        db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    obl2 = [o for o in rq.list_obligations(lib, db_path=_db(tmp_path))
            if o["obligation_id"] != obl["obligation_id"]][0]
    assert ms.get_status(obl2["obligation_id"], db_path=_db(tmp_path))["verdict"] == ms.V_UNCERTAIN

    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.confirm_obligation(rev["review_id"], obl["obligation_id"],
                            confirmed_by="复核人甲", reason="已人工核对原件，确属近三年口径",
                            db_path=_db(tmp_path))
    # ★ 确认后必须**重新批准**（批完再改是常见越权口子）
    assert pkg.get_review(rev["review_id"], db_path=_db(tmp_path))["status"] == pkg.ST_REVIEWED
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    out = Path(res["output_dir"])
    assert res["exported"] == 1, res
    assert res["gate_blocked_obligations"] == 1, res
    delivered = _delivered_files(out)
    assert len(delivered) == 1, delivered


def test_confirmation_requires_named_confirmer_and_reason(tmp_path):
    """人工确认**必须**可追责（人 + 理由）。"""
    lib, _root, obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    with pytest.raises(ValueError, match="confirmed_by"):
        pkg.confirm_obligation(rev["review_id"], obl["obligation_id"],
                                confirmed_by="  ", reason="r", db_path=_db(tmp_path))
    with pytest.raises(ValueError, match="reason"):
        pkg.confirm_obligation(rev["review_id"], obl["obligation_id"],
                                confirmed_by="复核人甲", reason="  ", db_path=_db(tmp_path))


def test_confirmation_is_never_promoted_to_a_global_rule(tmp_path):
    """一次确认**只**落 review_id 维度 ⇒ 不得升级为全局规则。"""
    lib, _root, obl = _uncertain_library(tmp_path)
    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    out = pkg.confirm_obligation(rev["review_id"], obl["obligation_id"],
                                 confirmed_by="复核人甲", reason="人工核对无误",
                                 db_path=_db(tmp_path))
    assert out["promoted_to_global_rule"] is False
    confs = pkg.list_confirmations(rev["review_id"], db_path=_db(tmp_path))
    assert len(confs) == 1
    assert confs[0]["confirmed_by"] == "复核人甲"
    assert confs[0]["reason"] == "人工核对无误"


def test_confirmation_cannot_release_restricted_material(tmp_path):
    """★ 三道闸**取其严**：人工确认**不能**让受限材料进交付包。"""
    lib, _root, obl = _uncertain_library(tmp_path)
    # 把材料改成**受限**（敏感等级未知 ⇒ 不可外发），并**重跑匹配**
    # （可交付资格是判定层的产物 ⇒ 改了卡必须重跑匹配才生效）
    conn = mf._connect(_db(tmp_path))
    conn.execute("UPDATE dd_v1_material_cards SET sensitivity=''")
    conn.commit()
    ms.match_library(lib, db_path=_db(tmp_path))

    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.confirm_obligation(rev["review_id"], obl["obligation_id"],
                            confirmed_by="复核人甲", reason="人工核对无误",
                            db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    out = Path(res["output_dir"])
    assert res["exported"] == 0, "人工确认不得绕过披露闸"
    blocked = _json.loads((out / "被拦截清单.json").read_text(encoding="utf-8"))
    assert any("披露" in str(b.get("category") or "")
               or "敏感等级" in str(b.get("category") or "") for b in blocked), blocked


def test_batch_approval_does_not_release_every_uncertain_obligation(tmp_path):
    """★ 缺陷本体：单据级「批量批准」**不等于**逐条放行。

    修前：``approve()`` 一次批准 ⇒ 全部义务无条件复制。
    修后：单据批准只是**单据级**放行；逐条义务仍各自过闸。
    """
    lib, _root, _obl = _uncertain_library(tmp_path)
    for i in range(3, 8):
        rq.ingest_checklist(lib, [{"item_no": str(i), "raw_text": "请提供近三年审计报告"}],
                            db_path=_db(tmp_path))
    ms.match_library(lib, db_path=_db(tmp_path))
    obl = rq.list_obligations(lib, db_path=_db(tmp_path))
    assert len(obl) >= 2, obl

    rev = pkg.build_review_pack(lib, db_path=_db(tmp_path))
    pkg.approve(rev["review_id"], approved_by="复核人甲", db_path=_db(tmp_path))
    res = pkg.package_approved(rev["review_id"], db_path=_db(tmp_path))

    out = Path(res["output_dir"])
    assert res["exported"] == 0, "批量批准把不确定的义务全放行了"
    assert res["gate_blocked_obligations"] >= 2, res
    assert not _delivered_files(out)
